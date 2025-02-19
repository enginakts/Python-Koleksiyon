import tkinter as tk
from tkinter import filedialog, ttk, messagebox
from docx import Document
from fpdf import FPDF
from PIL import Image
import os
from pdf2docx import Converter
from pdf2image import convert_from_path

class DosyaDonusturucu:
    def __init__(self):
        self.root = tk.Tk()
        self.root.title("Çoklu Format Dönüştürücü")
        self.root.geometry("500x300")
        
        # Arayüz elemanları
        self.kaynak_format_label = tk.Label(self.root, text="Kaynak Format:")
        self.kaynak_format_label.pack(pady=5)
        
        self.kaynak_format = ttk.Combobox(self.root, 
            values=[".docx", ".pdf", ".jpg", ".png"])
        self.kaynak_format.pack(pady=5)
        self.kaynak_format.set(".docx")
        
        self.hedef_format_label = tk.Label(self.root, text="Hedef Format:")
        self.hedef_format_label.pack(pady=5)
        
        self.hedef_format = ttk.Combobox(self.root, 
            values=[".pdf", ".docx", ".jpg", ".png"])
        self.hedef_format.pack(pady=5)
        self.hedef_format.set(".pdf")
        
        self.yukle_button = tk.Button(self.root, text="Dosya Seç ve Dönüştür", 
            command=self.dosya_sec)
        self.yukle_button.pack(pady=20)
        
        self.durum_label = tk.Label(self.root, text="")
        self.durum_label.pack(pady=10)

    def docx_to_pdf(self, docx_path, pdf_path):
        doc = Document(docx_path)
        pdf = FPDF()
        pdf.add_page()
        pdf.set_font("Arial", size=12)
        
        for para in doc.paragraphs:
            pdf.cell(200, 10, txt=para.text, ln=True)
        
        pdf.output(pdf_path)

    def pdf_to_docx(self, pdf_path, docx_path):
        cv = Converter(pdf_path)
        cv.convert(docx_path)
        cv.close()

    def pdf_to_image(self, pdf_path, image_path):
        images = convert_from_path(pdf_path)
        images[0].save(image_path)

    def image_to_pdf(self, image_path, pdf_path):
        image = Image.open(image_path)
        rgb_image = image.convert('RGB')
        rgb_image.save(pdf_path)

    def dosya_sec(self):
        kaynak = self.kaynak_format.get()
        hedef = self.hedef_format.get()
        
        dosya_tipleri = {
            ".docx": "Word files",
            ".pdf": "PDF files",
            ".jpg": "JPEG files",
            ".png": "PNG files"
        }
        
        file_path = filedialog.askopenfilename(
            filetypes=[(dosya_tipleri[kaynak], f"*{kaynak}")])
            
        if file_path:
            try:
                output_path = os.path.splitext(file_path)[0] + hedef
                
                # DOCX -> PDF
                if kaynak == ".docx" and hedef == ".pdf":
                    self.docx_to_pdf(file_path, output_path)
                
                # PDF -> DOCX
                elif kaynak == ".pdf" and hedef == ".docx":
                    self.pdf_to_docx(file_path, output_path)
                
                # PDF -> Image (JPG/PNG)
                elif kaynak == ".pdf" and hedef in [".jpg", ".png"]:
                    self.pdf_to_image(file_path, output_path)
                
                # Image -> PDF
                elif kaynak in [".jpg", ".png"] and hedef == ".pdf":
                    self.image_to_pdf(file_path, output_path)
                
                # Image -> Image
                elif kaynak in [".jpg", ".png"] and hedef in [".jpg", ".png"]:
                    image = Image.open(file_path)
                    image.save(output_path)
                
                else:
                    messagebox.showerror("Hata", 
                        "Bu dönüşüm şu anda desteklenmiyor!")
                    return
                
                self.durum_label.config(
                    text=f"Dosya başarıyla dönüştürüldü:\n{output_path}")
                
            except Exception as e:
                messagebox.showerror("Hata", f"Dönüştürme sırasında hata: {str(e)}")

    def baslat(self):
        self.root.mainloop()

if __name__ == "__main__":
    app = DosyaDonusturucu()
    app.baslat()
